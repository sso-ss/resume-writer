import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from to_docx import create_resume_doc
from to_docx_right_sidebar import build_layout_c
from to_docx_right_sidebar_refined import build_layout_refined
from to_docx_two_column import build_layout_b


RESUME = """# Test Designer

Portfolio: https://example.com | designer@example.com

## Summary
Product Designer with experience building accessible products.

## Experience
### Product Designer | Example Studio | Jan 2024 - Present
- Improved task completion by 20% through usability testing.

## Recognition
- **Design Award** - Accessible product, 2025
* **Interaction Award** - Research project, 2024

## Education
B.Des | Example University | 2020

## Skills & Tools
- **Design:** Interaction Design, Accessibility
- **Tools:** Figma, FigJam
"""


class LayoutTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / "resume.md"
        self.source.write_text(RESUME, encoding="utf-8")
        self.output = Path(self.directory.name) / "resume.docx"

    def build(self, builder):
        builder(str(self.source), str(self.output))
        return Document(self.output)

    def test_resume_text_respects_minimum_size(self):
        for builder in (create_resume_doc, build_layout_b, build_layout_c, build_layout_refined):
            with self.subTest(layout=builder.__name__):
                document = self.build(builder)
                self.assertEqual(document.styles['Normal'].font.size, Pt(10.5))
                for size in document.element.xpath('.//w:r[w:t]/w:rPr/w:sz'):
                    self.assertGreaterEqual(int(size.get(qn('w:val'))), 18)

    def test_left_grid_and_cells_use_one_to_two_ratio(self):
        table = self.build(build_layout_b).tables[0]
        expected_widths = [Inches(2.4), Inches(4.8)]
        self.assertFalse(table.autofit)
        self.assertEqual([column.width for column in table.columns], expected_widths)
        self.assertEqual([cell.width for cell in table.rows[0].cells], expected_widths)

    def test_left_education_uses_standard_section_spacing(self):
        sidebar = self.build(build_layout_b).tables[0].cell(0, 0)
        paragraphs = sidebar.paragraphs
        education_index = next(
            index for index, paragraph in enumerate(paragraphs)
            if paragraph.text == "EDUCATION"
        )
        self.assertEqual(paragraphs[education_index - 1].text, "Tools: Figma, FigJam")
        self.assertEqual(paragraphs[education_index].paragraph_format.space_before, Pt(10))
        self.assertEqual(paragraphs[education_index].paragraph_format.space_after, Pt(3))

    def test_left_without_education_has_no_trailing_separator(self):
        self.source.write_text(
            RESUME.replace("## Education\nB.Des | Example University | 2020\n\n", ""),
            encoding="utf-8",
        )
        sidebar = self.build(build_layout_b).tables[0].cell(0, 0)
        self.assertEqual(sidebar.paragraphs[-1].text, "Tools: Figma, FigJam")

    def test_refined_recognition_has_bullets_but_skills_and_education_do_not(self):
        document = self.build(build_layout_refined)
        sidebar = document.tables[0].cell(0, 1)
        awards = [paragraph for paragraph in sidebar.paragraphs if "Award -" in paragraph.text]
        self.assertEqual(len(awards), 2)
        for paragraph in awards:
            self.assertEqual(paragraph.style.name, "List Bullet")
            self.assertEqual(paragraph.paragraph_format.left_indent, Inches(0.15))
            self.assertEqual(paragraph.paragraph_format.first_line_indent, -Inches(0.1))
            self.assertEqual(paragraph.paragraph_format.tab_stops[0].position, Inches(0.15))
            self.assertTrue(paragraph.runs[0].bold)
        award_texts = {paragraph.text for paragraph in awards}
        for paragraph in sidebar.paragraphs:
            if paragraph.text not in award_texts:
                self.assertNotEqual(paragraph.style.name, "List Bullet")

    def test_main_bullets_have_explicit_hanging_indent_and_tab_stop(self):
        for builder, main_index, bullet_count in (
            (build_layout_b, 1, 3),
            (build_layout_c, 0, 3),
            (build_layout_refined, 0, 1),
        ):
            with self.subTest(layout=builder.__name__):
                document = self.build(builder)
                main = document.tables[-1].cell(0, main_index)
                bullets = [paragraph for paragraph in main.paragraphs if paragraph.style.name == "List Bullet"]
                self.assertEqual(len(bullets), bullet_count)
                for paragraph in bullets:
                    formatting = paragraph.paragraph_format
                    self.assertEqual(formatting.left_indent, Inches(0.15))
                    self.assertEqual(formatting.first_line_indent, -Inches(0.15))
                    self.assertEqual(formatting.tab_stops[0].position, Inches(0.15))


if __name__ == "__main__":
    unittest.main()
