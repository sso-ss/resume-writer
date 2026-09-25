#!/usr/bin/env python3
"""Generate an editable Word counterpart to the editorial HTML print layout."""

import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from lxml import html

from docx_shared import remove_table_borders, set_cell_margins
from to_html_editorial import build_html


BODY_FONT = "DM Sans"
DISPLAY_FONT = "Manrope"
TEXT = "232924"
ACCENT = "466454"
MAIN_PT = 374.6
GAP_PT = 24
SIDEBAR_PT = 132.75


def add_text(paragraph, text, size=7.5, bold=False, color=TEXT, font=BODY_FONT):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def inline_runs(paragraph, element, size=7.5, bold=False, color=TEXT):
    bold = bold or element.tag in ("strong", "b") or element.get("class") == "label"
    if element.get("class") == "org":
        size, bold = 7.875, False
    if element.text:
        add_text(paragraph, element.text, size, bold, color)
    for child in element:
        if child.tag == "br":
            paragraph.add_run().add_break()
        elif child.tag == "a" and child.get("href", "").startswith(("https://", "http://", "mailto:")):
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(
                child.get("href"), RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            ))
            run = add_text(paragraph, child.text_content(), size, bold, "365D47")
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        else:
            inline_runs(paragraph, child, size, bold, color)
        if child.get("class") == "label":
            paragraph.add_run().add_break()
        if child.tail:
            add_text(paragraph, child.tail, size, bold, color)


def paragraph(container, after=0, before=0, line=1.4, keep=False):
    if hasattr(container, "_tc") and len(container.paragraphs) == 1 and not container.paragraphs[0].text:
        result = container.paragraphs[0]
    else:
        result = container.add_paragraph()
    formatting = result.paragraph_format
    formatting.space_before = Pt(before)
    formatting.space_after = Pt(after)
    formatting.line_spacing = line
    formatting.keep_with_next = keep
    return result


def section_heading(container, text, first=False):
    result = paragraph(container, after=6, before=0 if first else 12, keep=True)
    run = add_text(result, text.upper(), 6.75, True, ACCENT)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "23")
    run._r.get_or_add_rPr().append(spacing)


def bullet_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_id = max(int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "\u2022"), ("suff", "tab")):
        item = OxmlElement(f"w:{tag}")
        item.set(qn("w:val"), value)
        level.append(item)
    abstract.append(level)
    numbering.append(abstract)
    return numbering.add_num(abstract_id).numId


def render_main(cell, element, numbering_id):
    for section_index, section in enumerate(element.findall("section")):
        section_heading(cell, section.find("h2").text_content(), first=section_index == 0)
        for article in section.findall("article"):
            if article.get("class") == "job":
                header = article.find("div")
                result = paragraph(cell, after=4.5, line=1.3, keep=True)
                result.paragraph_format.tab_stops.add_tab_stop(Pt(MAIN_PT), WD_TAB_ALIGNMENT.RIGHT)
                inline_runs(result, header.find("h3"), 9.375, True)
                add_text(result, "\t")
                inline_runs(result, header.find("p"), 7.125, color="747C76")
                bullets = article.findall("ul/li")
                for index, bullet in enumerate(bullets):
                    last = index == len(bullets) - 1
                    result = paragraph(cell, after=8.25 if last else 2.25, keep=not last)
                    formatting = result.paragraph_format
                    formatting.left_indent = Pt(12)
                    formatting.first_line_indent = Pt(-10.5)
                    formatting.tab_stops.add_tab_stop(Pt(12))
                    num_properties = result._p.get_or_add_pPr().get_or_add_numPr()
                    num_properties.get_or_add_ilvl().val = 0
                    num_properties.get_or_add_numId().val = numbering_id
                    inline_runs(result, bullet, 7.65)
            else:
                result = paragraph(cell, after=2.25, line=1.3, keep=True)
                inline_runs(result, article.find("h3"), 9.375, True)
                details = article.findall("p")
                for index, detail in enumerate(details):
                    last = index == len(details) - 1
                    result = paragraph(cell, after=5.25 if last else 1.5, keep=not last)
                    inline_runs(result, detail, color="365D47" if detail.get("class") == "portfolio" else TEXT)


def render_sidebar(cell, element):
    for section_index, section in enumerate(element.findall("section")):
        section_heading(cell, section.find("h2").text_content(), first=section_index == 0)
        details = section.findall("p")
        for index, detail in enumerate(details):
            result = paragraph(cell, after=6, line=1.45, keep=index < len(details) - 1)
            inline_runs(result, detail)


def build_layout_editorial(md_path, docx_path=None, *, editorial_header=None):
    source = Path(md_path)
    if docx_path is None:
        base = source.stem.replace("_Resume", "")
        docx_path = source.with_name(f"{base}_ProductDesigner_Resume_Editorial.docx")
    with tempfile.TemporaryDirectory(prefix="editorial-word-") as directory:
        html_path = build_html(str(source), str(Path(directory) / "resume.html"))
        tree = html.parse(html_path)
    page = tree.xpath('//main[@class="page"]')[0]
    header = page.find("header")
    if editorial_header:
        for class_name in ("eyebrow", "role"):
            if class_name in editorial_header:
                header.xpath(f'./p[@class="{class_name}"]')[0].text = editorial_header[class_name]
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.4), Inches(.34)
    section.left_margin = section.right_margin = Inches(.56)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = BODY_FONT, Pt(7.5)
    normal.paragraph_format.space_after = Pt(0)
    for class_name, size, after, color in (("eyebrow", 6.75, 6.75, ACCENT),):
        result = paragraph(document, after=after)
        text = header.xpath(f'./p[@class="{class_name}"]')[0].text_content()
        run = add_text(result, text.upper(), size, True, color)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "30")
        run._r.get_or_add_rPr().append(spacing)
    result = paragraph(document, after=5.25, line=1.1)
    add_text(result, header.find("h1").text_content(), 29.5, True, "1D2821", DISPLAY_FONT)
    result = paragraph(document, after=9.75)
    add_text(result, header.xpath('./p[@class="role"]')[0].text_content(), 12.75, color=ACCENT)
    result = paragraph(document, after=9.75)
    for index, contact in enumerate(header.find("div")):
        if index:
            add_text(result, "    ", 7.5)
        inline_runs(result, contact, color="58605A")
    result = paragraph(document, after=10.5, line=1.55)
    inline_runs(result, header.xpath('./p[@class="intro"]')[0], 8.625, color="454E47")
    layout = page.xpath('./div[@class="layout"]')[0]
    table = document.add_table(rows=1, cols=3)
    table.autofit = False
    for column, cell, width in zip(table.columns, table.rows[0].cells, (MAIN_PT, GAP_PT, SIDEBAR_PT)):
        column.width = cell.width = Pt(width)
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    remove_table_borders(table)
    main, gutter, sidebar = table.rows[0].cells
    render_main(main, layout.find("div"), bullet_numbering(document))
    render_sidebar(sidebar, layout.find("aside"))
    output = Path(docx_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return str(output)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit("Usage: to_docx_editorial.py input.md [output.docx]")
    print(build_layout_editorial(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None))