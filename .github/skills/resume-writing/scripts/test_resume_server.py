import io
import json
import re
import tempfile
import threading
import unittest
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document
from docx.oxml.ns import qn
from pypdf.generic import ContentStream
from pypdf import PdfReader
from playwright.sync_api import sync_playwright
from serve_resume import WORD_BUILDERS, create_server
from to_docx_right_sidebar_refined import build_layout_refined
from to_html_editorial import build_html


class ResumeServerTests(unittest.TestCase):
    def test_template_switch_preserves_edits_and_selected_word_export(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "resume.md"
            source.write_text(
                "# Test Designer\n\nPortfolio: https://example.com | test@example.com\n\n"
                "## Summary\n\nDesigner\n\n## Experience\n\n"
                "### Designer | Example | 2024 - Present\n- Built product\n\n"
                "## Education\n\nDesign degree\n\n## Skills & Tools\n\n- **Design:** Prototyping\n"
            )
            preview = build_html(str(source), str(Path(directory) / "resume.html"), "two-column-left")
            server = create_server(preview, "two-column-left")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch()
                    try:
                        page = browser.new_page()
                        page.route("https://**/*", lambda route: route.abort())
                        page.goto(f"http://127.0.0.1:{server.server_port}/")
                        self.assertEqual(page.locator("body").get_attribute("data-layout"), "two-column-left")
                        page.locator("h1").fill("Edited Designer")
                        background_before = page.locator("body").evaluate("node => getComputedStyle(node).backgroundColor")
                        page.locator("#accent-color").fill("#b02a36")
                        self.assertEqual(page.locator("h2").first.evaluate("node => getComputedStyle(node).color"), "rgb(176, 42, 54)")
                        self.assertEqual(page.locator("#save-word").evaluate("node => getComputedStyle(node).backgroundColor"), "rgb(176, 42, 54)")
                        self.assertEqual(page.locator("#save-word").evaluate("node => getComputedStyle(node).color"), "rgb(255, 255, 255)")
                        background_red = page.locator("body").evaluate("node => getComputedStyle(node).backgroundColor")
                        self.assertNotEqual(background_red, background_before)
                        page.locator("#accent-color").fill("#f5d948")
                        self.assertEqual(page.locator("#save-pdf").evaluate("node => getComputedStyle(node).backgroundColor"), "rgb(245, 217, 72)")
                        self.assertEqual(page.locator("#save-pdf").evaluate("node => getComputedStyle(node).color"), "rgb(17, 17, 17)")
                        self.assertNotEqual(page.locator("body").evaluate("node => getComputedStyle(node).backgroundColor"), background_red)
                        page.locator("#accent-color").fill("#b02a36")
                        for layout in WORD_BUILDERS:
                            page.locator("#preview-layout").select_option(layout)
                            self.assertEqual(page.locator("body").get_attribute("data-layout"), layout)
                            self.assertEqual(page.locator("h1").inner_text(), "Edited Designer")
                            self.assertEqual(page.locator(".page .contact").count(), 1)
                            self.assertEqual(page.locator(".eyebrow").is_visible(), layout == "editorial-html")
                            self.assertEqual(page.locator(".role").is_visible(), layout == "editorial-html")
                            self.assertEqual(page.locator(".summary-label").is_visible(), layout == "two-column-right-refined")
                            if layout == "two-column-right-refined":
                                self.assertEqual(page.locator(".contact").evaluate("node => getComputedStyle(node).borderBottomStyle"), "solid")
                                self.assertEqual(page.locator(".layout").evaluate("node => getComputedStyle(node).display"), "grid")
                        page.set_viewport_size({"width": 390, "height": 844})
                        for layout in WORD_BUILDERS:
                            page.locator("#preview-layout").select_option(layout)
                            self.assertGreater(page.locator(".main-column").bounding_box()["width"], 250)
                            self.assertFalse(page.evaluate("document.documentElement.scrollWidth > document.documentElement.clientWidth"))
                        page.locator("#preview-layout").select_option("single-column")
                        page.reload()
                        self.assertEqual(page.locator("#preview-layout").input_value(), "single-column")
                        self.assertEqual(page.locator("h1").inner_text(), "Edited Designer")
                        self.assertEqual(page.locator("#accent-color").input_value(), "#b02a36")
                        self.assertEqual(page.locator("#save-word").evaluate("node => getComputedStyle(node).backgroundColor"), "rgb(176, 42, 54)")
                        self.assertEqual(page.locator("body").evaluate("node => getComputedStyle(node).backgroundColor"), background_red)
                        self.assertEqual(page.locator(".page").evaluate("node => getComputedStyle(node).fontFamily"), "Arial, sans-serif")
                        with page.expect_download() as result:
                            page.get_by_role("button", name="Save as Word").click()
                        output = Path(directory) / "switched.docx"
                        result.value.save_as(output)
                        document = Document(output)
                        self.assertFalse(document.tables)
                        self.assertIn("Edited Designer", " ".join(p.text for p in document.paragraphs))
                        self.assertEqual(document.styles["Normal"].font.name, "Arial")
                        self.assertTrue(any(node.get(qn("w:val")) == "B02A36" for node in document.element.xpath(".//w:color")))
                    finally:
                        browser.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join()

    def test_pdf_button_download_preserves_browser_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "resume.md"
            source.write_text("# Original Name\n\nname@example.com\n\n## Summary\n\nOriginal summary\n")
            html = build_html(str(source), str(Path(directory) / "resume.html"))
            server = create_server(html, "editorial-html")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch()
                    try:
                        page = browser.new_page()
                        page.route("https://**/*", lambda route: route.abort())
                        page.goto(f"http://127.0.0.1:{server.server_port}/")
                        page.locator("h1").fill("Browser Edited Name")
                        page.locator("#preview-layout").select_option("two-column-right-refined")
                        with page.expect_download() as result:
                            page.get_by_role("button", name="Save as PDF").click()
                        download = result.value
                        self.assertEqual(download.suggested_filename, "Browser_Edited_Name_Resume.pdf")
                        output = Path(directory) / "download.pdf"
                        download.save_as(output)
                        pdf = PdfReader(output)
                        text = pdf.pages[0].extract_text()
                        self.assertIn("Browser Edited Name", text)
                        self.assertIn("SUMMARY", text.upper())
                        self.assertNotIn("Research", text)
                        self.assertEqual(len(pdf.pages), 1)
                    finally:
                        browser.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join()

    def test_all_document_layouts_have_word_generators(self):
        self.assertEqual(
            set(WORD_BUILDERS),
            {"single-column", "two-column-left", "two-column-right", "two-column-right-refined", "editorial-html"},
        )
        with self.assertRaisesRegex(ValueError, "Unknown resume layout"):
            create_server("missing.html", "unknown")

    def test_standalone_refined_layout_keeps_header_divider(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "resume.md"
            output = Path(directory) / "resume.docx"
            source.write_text("# Name\n\nname@example.com\n\n## Summary\n\nSummary\n")
            build_layout_refined(str(source), str(output))
            document = Document(output)
            self.assertEqual(len(document.element.xpath(".//w:pBdr/w:bottom")), 1)

    def test_export_uses_python_generator_and_rejects_foreign_requests(self):
        with tempfile.TemporaryDirectory() as directory:
            html = Path(directory) / "resume.html"
            html.write_text("<head></head><body>Resume</body>")
            server = create_server(html, "two-column-right-refined")
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                origin = f"http://127.0.0.1:{server.server_port}"
                preview = urlopen(origin).read().decode()
                token = re.search(r'content="([^"]+)"', preview).group(1)
                markdown = (
                    "# Edited Name\n\nPortfolio: https://example.com | test@example.com\n\n"
                    "## Summary\n\nEdited summary\n\n## Experience\n\n"
                    "### Designer | Edited Company | 2024 - Present\n- Edited achievement\n\n"
                    "## Education\n\nEdited degree\n"
                )
                body = json.dumps({"markdown": markdown}).encode()
                headers = {"Content-Type": "application/json", "Origin": origin, "X-Resume-Token": token}
                request = Request(origin + "/export-word", data=body, headers=headers)
                document = Document(io.BytesIO(urlopen(request).read()))
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                text += "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                for expected in ("Edited Name", "Edited Company", "Edited achievement", "Edited degree"):
                    self.assertIn(expected, text)
                self.assertEqual(len(document.tables[0].columns), 2)
                self.assertEqual(len(document.element.xpath(".//w:pBdr/w:bottom")), 1)
                with self.assertRaises(HTTPError) as error:
                    urlopen(Request(origin + "/check-ats", data=body, headers=headers))
                self.assertEqual(error.exception.code, 404)
                selected_body = json.dumps({"markdown": markdown, "layout": "single-column"}).encode()
                selected = Document(io.BytesIO(urlopen(Request(origin + "/export-word", data=selected_body, headers=headers)).read()))
                self.assertFalse(selected.tables)
                invalid_body = json.dumps({"markdown": markdown, "layout": "unknown"}).encode()
                with self.assertRaises(HTTPError) as error:
                    urlopen(Request(origin + "/export-word", data=invalid_body, headers=headers))
                self.assertEqual(error.exception.code, 400)
                invalid_color = json.dumps({"markdown": markdown, "accent_color": "red; color: black"}).encode()
                with self.assertRaises(HTTPError) as error:
                    urlopen(Request(origin + "/export-word", data=invalid_color, headers=headers))
                self.assertEqual(error.exception.code, 400)
                left_body = json.dumps({"markdown": markdown, "layout": "two-column-left", "accent_color": "#b02a36"}).encode()
                left_doc = Document(io.BytesIO(urlopen(Request(origin + "/export-word", data=left_body, headers=headers)).read()))
                sidebar = left_doc.tables[0].cell(0, 0)
                self.assertEqual(sidebar._tc.xpath("./w:tcPr/w:shd")[0].get(qn("w:fill")), "F0EFEB")
                self.assertEqual(left_doc.styles["Normal"].font.name, "Arial")
                self.assertTrue(any(node.get(qn("w:val")) == "B02A36" for node in left_doc.element.xpath(".//w:color")))
                left_pdf = PdfReader(io.BytesIO(urlopen(Request(origin + "/export-pdf", data=left_body, headers=headers)).read()))
                self.assertIn("Edited Name", " ".join(left_pdf.pages[0].extract_text().split()))
                stream = ContentStream(left_pdf.pages[0].get_contents(), left_pdf)
                self.assertTrue(any(
                    operator == b"rg" and all(abs(float(value) - channel / 255) < .01 for value, channel in zip(operands, (240, 239, 235)))
                    for operands, operator in stream.operations
                ))
                self.assertTrue(any(
                    operator == b"rg" and all(abs(float(value) - channel / 255) < .01 for value, channel in zip(operands, (176, 42, 54)))
                    for operands, operator in stream.operations
                ))
                editorial_body = json.dumps({"markdown": markdown, "layout": "editorial-html", "accent_color": "#b02a36"}).encode()
                editorial_doc = Document(io.BytesIO(urlopen(Request(origin + "/export-word", data=editorial_body, headers=headers)).read()))
                self.assertTrue(any(node.get(qn("w:val")) == "B02A36" for node in editorial_doc.element.xpath(".//w:color")))
                response = urlopen(Request(origin + "/export-pdf", data=body, headers=headers))
                self.assertEqual(response.headers["Content-Type"], "application/pdf")
                pdf = PdfReader(io.BytesIO(response.read()))
                pdf_text = " ".join(" ".join(page.extract_text() or "" for page in pdf.pages).split())
                for expected in ("Edited Name", "Edited Company", "Edited achievement", "Edited degree", "EXPERIENCE", "EDUCATION"):
                    self.assertIn(expected, pdf_text)
                self.assertLess(pdf_text.index("Edited achievement"), pdf_text.index("Edited degree"))
                self.assertEqual(len(pdf.pages), 1)
                self.assertIsNotNone(pdf.trailer["/Root"].get("/StructTreeRoot"))
                for bad_headers in ({}, {**headers, "Origin": "https://example.com"}):
                    with self.assertRaises(HTTPError) as error:
                        urlopen(Request(origin + "/export-word", data=body, headers=bad_headers))
                    self.assertEqual(error.exception.code, 403)
                    with self.assertRaises(HTTPError) as error:
                        urlopen(Request(origin + "/export-pdf", data=body, headers=bad_headers))
                    self.assertEqual(error.exception.code, 403)
                with self.assertRaises(HTTPError) as error:
                    urlopen(origin + "/../Jennifer_Lauren_Resume.md")
                self.assertEqual(error.exception.code, 404)
            finally:
                server.shutdown()
                server.server_close()
                thread.join()


if __name__ == "__main__":
    unittest.main()