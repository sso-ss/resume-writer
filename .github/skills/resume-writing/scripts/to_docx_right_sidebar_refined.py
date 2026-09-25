#!/usr/bin/env python3
"""Convert a Markdown resume to the refined two-column right-sidebar layout.

Usage: python to_docx_right_sidebar_refined.py <input.md> [output.docx]
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Optional

try:
    from docx import Document
    from docx.shared import Inches, Pt
except ImportError:
    print("Error: python-docx is required. Install it with: pip install python-docx")
    sys.exit(1)

from docx_shared import (
    ACCENT, FONT, MAIN_W, SIDEBAR_W, TEXT_BODY, TEXT_DARK, TEXT_MUTED,
    add_bottom_rule, fmt, heading, parse_contact_items, parse_markdown,
    remove_table_borders, render_main_content, render_sidebar_skills,
    set_cell_margins, sidebar_lines,
)


def build_layout_refined(
    md_path: str, docx_path: Optional[str] = None, *, show_header_divider: bool = True
) -> str:
    md_file = Path(md_path)
    if not md_file.exists():
        print(f"Error: {md_path} not found")
        sys.exit(1)

    if docx_path is None:
        base = md_file.stem.replace("_Resume", "")
        docx_path = str(md_file.with_name(f"{base}_ProductDesigner_Resume_RightRefined.docx"))

    name, contact_line, sections = parse_markdown(md_file.read_text(encoding="utf-8"))
    contact_items = parse_contact_items(contact_line) if contact_line else []

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.4)
    sec.bottom_margin = Inches(0.4)
    sec.left_margin = Inches(0.4)
    sec.right_margin = Inches(0.4)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)
    style.font.color.rgb = TEXT_BODY
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(name or "Full Name")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = TEXT_DARK
    r.font.name = FONT

    if contact_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        for index, (label, value) in enumerate(contact_items):
            if index:
                p.add_run("  |  ")
            r = p.add_run(f"{label}: {value}" if label.lower() == "portfolio" else value)
            r.font.size = Pt(9)
            r.font.color.rgb = ACCENT if label.lower() == "portfolio" else TEXT_MUTED
            r.font.name = FONT

    if show_header_divider:
        rule = doc.add_paragraph()
        rule.paragraph_format.space_before = Pt(4)
        rule.paragraph_format.space_after = Pt(4)
        add_bottom_rule(rule)

    if "summary" in sections:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("SUMMARY")
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = ACCENT
        for line in sections["summary"]:
            if line.strip():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                fmt(p, line.strip(), size=Pt(9.5))

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = MAIN_W
    table.columns[1].width = SIDEBAR_W
    main, sidebar = table.rows[0].cells
    main.width = MAIN_W
    sidebar.width = SIDEBAR_W

    set_cell_margins(sidebar, top=120, start=140, bottom=120, end=100)
    set_cell_margins(main, top=80, start=60, bottom=80, end=160)

    render_sidebar_skills(sidebar, sections)
    if "recognition" in sections:
        heading(sidebar, "Recognition")
        sidebar_lines(sidebar, sections["recognition"], bullets=True)
    if "education" in sections:
        heading(sidebar, "Education")
        sidebar_lines(sidebar, sections["education"])

    render_main_content(main, {
        key: value for key, value in sections.items()
        if key not in ("summary", "recognition")
    })

    for paragraph in [*doc.paragraphs, *main.paragraphs, *sidebar.paragraphs]:
        for run in paragraph.runs:
            if run.font.size is not None and run.font.size < Pt(9):
                run.font.size = Pt(9)

    remove_table_borders(table)
    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python to_docx_right_sidebar_refined.py <input.md> [output.docx]")
        sys.exit(1)

    result = build_layout_refined(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
    print(f"Refined right-sidebar resume saved to: {result}")