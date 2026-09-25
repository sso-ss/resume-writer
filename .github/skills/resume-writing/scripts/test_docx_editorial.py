import io
import tempfile
import threading
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Pt
from playwright.sync_api import sync_playwright

from serve_resume import create_server
from to_docx_editorial import build_layout_editorial
from to_html_editorial import build_html


SOURCE = """# Test Designer

Portfolio: https://example.com | designer@example.com

## Summary
Senior Product Designer with 6+ years creating accessible tools.

## Experience
### Product Designer | ExampleCo | Jan 2024 - Present
- Improved task completion by 20%.

## Key Projects
- **Portal** - Accessible tools. Case study: https://example.com/portal

## Skills & Tools
- **Design:** Interaction, Accessibility
- **Tools:** Figma

## Recognition
- **Design Award** - Portal, 2025

## Education
B.Des | Example University | 2020
"""


class EditorialWordTests(unittest.TestCase):
    def test_preview_button_exports_editorial_and_preserves_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'resume.md'
            source.write_text(SOURCE)
            preview = build_html(str(source), str(Path(directory) / 'resume.html'))
            server = create_server(preview)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                with sync_playwright() as playwright:
                    browser = playwright.chromium.launch()
                    try:
                        page = browser.new_page()
                        page.route('https://**/*', lambda route: route.abort())
                        page.goto(f'http://127.0.0.1:{server.server_port}/')
                        page.locator('h1').fill('Browser Edited Name')
                        page.locator('.role').fill('Staff Designer')
                        page.locator('.eyebrow').fill('Design and Research')
                        page.locator('.job .org').fill('Edited Company')
                        with page.expect_download() as result:
                            page.get_by_role('button', name='Save as Word', exact=True).click()
                        output = Path(directory) / 'browser.docx'
                        result.value.save_as(output)
                        document = Document(output)
                        text = '\n'.join(item.text for item in document.paragraphs)
                        self.assertIn('Browser Edited Name', text)
                        self.assertIn('Staff Designer', text)
                        self.assertIn('DESIGN AND RESEARCH', text)
                        self.assertNotIn('SUMMARY', text)
                        main, gutter, sidebar = document.tables[0].rows[0].cells
                        self.assertIn('Edited Company', main.text)
                        self.assertIn('SELECTED PROJECTS', main.text)
                        self.assertIn('EXPERTISE', sidebar.text)
                        self.assertEqual(document.element.xpath('.//w:pBdr'), [])
                    finally:
                        browser.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join()

    def test_editorial_structure_typography_and_dates(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "resume.md"
            source.write_text(SOURCE)
            output = build_layout_editorial(source, Path(directory) / "resume.docx")
            document = Document(output)
            header = '\n'.join(item.text for item in document.paragraphs)
            self.assertIn('RESEARCH', header)
            self.assertIn('Senior Product Designer', header)
            self.assertNotIn('SUMMARY', header)
            self.assertEqual(document.element.xpath('.//w:pBdr'), [])
            name = next(item for item in document.paragraphs if item.text == 'Test Designer')
            self.assertEqual(name.runs[0].font.name, 'Manrope')
            self.assertEqual(name.runs[0].font.size, Pt(29.5))
            main, gutter, sidebar = document.tables[0].rows[0].cells
            self.assertEqual([column.width for column in document.tables[0].columns],
                             [Pt(374.6), Pt(24), Pt(132.75)])
            self.assertIn('SELECTED PROJECTS', main.text)
            self.assertIn('EXPERTISE', sidebar.text)
            self.assertIn('TOOLS', sidebar.text)
            self.assertIn('RECOGNITION', sidebar.text)
            self.assertIn('EDUCATION', sidebar.text)
            job = next(item for item in main.paragraphs if 'ExampleCo' in item.text)
            self.assertEqual(job.text, 'Product Designer | ExampleCo\tJan 2024 - Present')
            self.assertEqual(job.paragraph_format.tab_stops[0].alignment, WD_TAB_ALIGNMENT.RIGHT)
            self.assertEqual(len(main._tc.xpath('.//w:numPr')), 1)
            self.assertEqual(len(sidebar._tc.xpath('.//w:numPr')), 0)
            self.assertEqual(len(document.element.xpath('.//w:hyperlink')), 3)

    def test_edited_editorial_header_is_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'resume.md'
            source.write_text(SOURCE)
            document = Document(build_layout_editorial(source, editorial_header={
                'eyebrow': 'Design and Research', 'role': 'Staff Designer',
            }))
            text = '\n'.join(item.text for item in document.paragraphs)
            self.assertIn('DESIGN AND RESEARCH', text)
            self.assertIn('Staff Designer', text)


if __name__ == '__main__':
    unittest.main()