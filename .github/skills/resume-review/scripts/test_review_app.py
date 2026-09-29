import base64
import json
import re
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from unittest.mock import patch

from docx import Document
from render_review import Block, match_findings, match_bullet_reviews, match_job_requirements
from review_app import SKILL_DIR, ReviewApplication, ReviewError, create_app_server, run_codex, validate_upload, expand_review, run_claude, detect_provider, provider_config, run_cursor, run_copilot, provider_executable

BULLET = 'Designed dashboards for users.'
RESUME = '# Test Candidate\n\n## Experience\n\n### Designer | Example | 2024\n- ' + BULLET + '\n'
SOURCE = 'Product Designer at Example. Design analytics dashboards, conduct usability research, and collaborate with engineers.'


def valid_review():
    review = {
        'candidate': 'Test Candidate', 'overall': 'Relevant dashboard experience; clarify outcomes.',
        'strengths': ['Dashboard experience'],
        'target_job': {'title': 'Product Designer', 'company': 'Example',
                       'source': 'User-provided job description', 'summary': 'Design analytics dashboards.',
                       'requirements': [{'requirement': 'Design analytics dashboards', 'priority': 'core',
                                         'status': 'partial-match', 'evidence_quotes': [BULLET],
                                         'feedback': 'Clarify the scope and outcomes.'}]},
        'bullet_reviews': [{'section': 'Experience', 'quote': BULLET, 'rating': 'needs-work',
                            'assessment': 'Scope and outcomes are missing.', 'strengths': ['Clear action'],
                            'gaps': ['No outcome'], 'suggestion': 'Add a verified outcome.',
                            'job_alignment': 'partial-match', 'job_feedback': 'Relevant to dashboards.'}],
        'findings': [{'section': 'Experience', 'severity': 'important', 'title': 'Clarify outcomes',
                      'quote': BULLET, 'issue': 'No outcome is described.', 'why': 'Impact is unclear.',
                      'suggestion': 'Add a verified outcome.'}],
    }

    del review['candidate']
    del review['target_job']['source']
    for requirement in review['target_job']['requirements']:
        del requirement['evidence_quotes']
        requirement['evidence_ids'] = ['r4']
    for entry in review['bullet_reviews'] + review['findings']:
        del entry['section']
        del entry['quote']
        entry['block_id'] = 'r4'
    for finding in review['findings']:
        finding['excerpt'] = ''
    return review


def upload(**changes):
    payload = {'filename': 'resume.md', 'file': base64.b64encode(RESUME.encode()).decode(), 'source': SOURCE, 'provider': 'codex'}
    payload.update(changes)
    return payload


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def wait_job(self, app, identifier):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            status = json.loads((app.directory(identifier) / 'status.json').read_text())
            if status['status'] in ('failed', 'complete') and not app.busy:
                return status
            time.sleep(.01)
        self.fail('Review did not finish')

    def test_upload_to_verified_review_loads_skill_and_preserves_original(self):
        prompts = []
        def runner(prompt, directory):
            prompts.append(prompt)
            return valid_review()
        app = ReviewApplication(self.root, runner)
        identifier = app.submit(upload(filename='../../elsewhere.md'))
        status = self.wait_job(app, identifier)
        self.assertEqual('complete', status['status'])
        self.assertEqual(RESUME, (app.directory(identifier) / 'resume.md').read_text())
        self.assertFalse((self.root.parent / 'elsewhere.md').exists())
        self.assertIn((SKILL_DIR / 'references/review-rubric.md').read_text(encoding='utf-8'), prompts[0])
        self.assertNotIn((SKILL_DIR / 'SKILL.md').read_text(encoding='utf-8'), prompts[0])
        guidelines = SKILL_DIR.parent / 'resume-writing/references/recruiter-guidelines.md'
        self.assertNotIn(guidelines.read_text(encoding='utf-8'), prompts[0])
        html = (app.directory(identifier) / 'review.html').read_text()
        self.assertIn('data-notes="note-1"', html)
        self.assertNotIn('Anchor not found', html)
        metrics = json.loads((app.directory(identifier) / 'metrics.json').read_text())
        self.assertEqual(1, metrics['attempts'])
        self.assertGreater(metrics['prompt_chars'], 0)
        self.assertGreater(metrics['response_chars'], 0)
        self.assertGreaterEqual(metrics['total_seconds'], metrics['analysis_seconds'])

    def test_bad_evidence_is_repaired_before_publishing(self):
        calls = []
        def runner(prompt, directory):
            calls.append(prompt)
            review = valid_review()
            if len(calls) == 1:
                review['findings'][0]['block_id'] = 'r999'
            return review
        app = ReviewApplication(self.root, runner)
        status = self.wait_job(app, app.submit(upload()))
        self.assertEqual('complete', status['status'])
        self.assertEqual(2, len(calls))
        self.assertIn('Validation error', calls[1])

    def test_persistent_bad_evidence_is_not_published(self):
        review = valid_review()
        review['bullet_reviews'] = []
        app = ReviewApplication(self.root, lambda *args: review)
        identifier = app.submit(upload())
        self.assertEqual('failed', self.wait_job(app, identifier)['status'])
        self.assertFalse((app.directory(identifier) / 'review.html').exists())

    def test_references_preserve_duplicate_text_and_exact_source(self):
        blocks = [Block('Test Candidate', 'name', 'Header'),
                  Block('Experience', 'section', 'Experience'),
                  Block(BULLET, 'bullet', 'Experience'),
                  Block(BULLET, 'bullet', 'Experience')]
        compact = valid_review()
        first = dict(compact['bullet_reviews'][0], block_id='r3')
        compact['bullet_reviews'].append(first)  # Deliberately reverse identical bullets.
        review = expand_review(compact, blocks, 'https://example.com/job', 'url')
        self.assertEqual('https://example.com/job', review['target_job']['source'])
        self.assertEqual({3: 0, 2: 1}, match_bullet_reviews(blocks, review['bullet_reviews']))
        self.assertEqual(({3: [0]}, []), match_findings(blocks, review['findings']))
        self.assertEqual({3: [0]}, match_job_requirements(blocks, review['target_job']['requirements']))
        self.assertNotIn('quote', compact['bullet_reviews'][0])

    def test_invalid_references_excerpts_and_coverage_are_rejected(self):
        blocks = [Block('Test Candidate', 'name', 'Header'),
                  Block('Experience', 'section', 'Experience'),
                  Block('Role', 'role', 'Experience'), Block(BULLET, 'bullet', 'Experience')]
        cases = []
        review = valid_review()
        review['target_job']['requirements'][0]['evidence_ids'] = ['r999']
        cases.append(review)
        review = valid_review()
        review['findings'][0]['excerpt'] = 'Invented achievement'
        cases.append(review)
        review = valid_review()
        review['bullet_reviews'].append(dict(review['bullet_reviews'][0]))
        cases.append(review)
        review = valid_review()
        review['bullet_reviews'][0]['block_id'] = 'r3'
        cases.append(review)
        for review in cases:
            with self.subTest(review=review), self.assertRaises(ValueError):
                expand_review(review, blocks, SOURCE, 'description')

    def test_partial_excerpt_and_no_reviewable_bullets(self):
        blocks = [Block('Test Candidate', 'name', 'Header'),
                  Block(BULLET, 'paragraph', 'Summary')]
        compact = valid_review()
        compact['bullet_reviews'] = []
        compact['findings'][0].update(block_id='r2', excerpt='dashboards')
        compact['target_job']['requirements'][0]['evidence_ids'] = ['r2']
        review = expand_review(compact, blocks, SOURCE, 'description')
        self.assertEqual('dashboards', review['findings'][0]['quote'])
        self.assertEqual({}, match_bullet_reviews(blocks, review['bullet_reviews']))
        self.assertEqual('Summary', review['findings'][0]['section'])

    def test_provider_failure_releases_queue_and_preserves_upload(self):
        def runner(*args):
            raise ReviewError('Please sign in to Codex.')
        app = ReviewApplication(self.root, runner)
        identifier = app.submit(upload())
        self.assertIn('sign in', self.wait_job(app, identifier)['message'])
        self.assertFalse(app.busy)
        self.assertTrue((app.directory(identifier) / 'resume.md').exists())

    def test_validation_rejects_missing_job_unsupported_corrupt_empty_and_large_files(self):
        app = ReviewApplication(self.root, lambda *args: self.fail('Invalid upload reached analysis'))
        for payload in [upload(source=''), upload(filename='resume.txt'), upload(file='@@'),
                        upload(file=''), upload(filename='resume.docx'),
                        upload(file=base64.b64encode(b'a' * (5 * 1024 * 1024 + 1)).decode())]:
            with self.subTest(payload_keys=list(payload)):
                with self.assertRaises(ReviewError):
                    app.submit(payload)
                self.assertFalse(app.busy)
        self.assertEqual([], list(self.root.iterdir()))

    def test_word_upload_extracts_text(self):
        source = self.root / 'sample.docx'
        doc = Document()
        doc.add_heading('Test Candidate', 0)
        doc.add_heading('Experience', 1)
        doc.add_paragraph(BULLET, style='List Bullet')
        doc.save(source)
        target = self.root / 'upload'
        target.mkdir()
        path, blocks = validate_upload(upload(filename='resume.docx', file=base64.b64encode(source.read_bytes()).decode()), target)
        self.assertTrue(any(block.text == BULLET and block.kind == 'bullet' for block in blocks))

    def test_busy_job_is_not_duplicated(self):
        entered, release = threading.Event(), threading.Event()
        def runner(*args):
            entered.set()
            release.wait(3)
            return valid_review()
        app = ReviewApplication(self.root, runner)
        identifier = app.submit(upload())
        try:
            self.assertTrue(entered.wait(2))
            with self.assertRaisesRegex(ReviewError, 'already running'):
                app.submit(upload())
        finally:
            release.set()
            self.wait_job(app, identifier)

    def test_restart_marks_interrupted_jobs_as_failed(self):
        directory = self.root / ('a' * 32)
        directory.mkdir()
        (directory / 'status.json').write_text(json.dumps({'status': 'reviewing'}))
        app = ReviewApplication(self.root)
        self.assertEqual('failed', json.loads((directory / 'status.json').read_text())['status'])

    def test_cli_uses_skill_prompt_read_only_and_structured_output(self):
        def fake_run(command, **kwargs):
            self.assertIn('--ephemeral', command)
            self.assertEqual('read-only', command[command.index('--sandbox') + 1])
            self.assertEqual('skill prompt', kwargs['input'])
            Path(command[command.index('--output-last-message') + 1]).write_text(json.dumps({'review': valid_review(), 'error': ''}))
            return type('Result', (), {'returncode': 0})()
        with patch('review_app.shutil.which', return_value='/bin/codex'), patch('review_app.subprocess.run', side_effect=fake_run):
            self.assertEqual(valid_review(), run_codex('skill prompt', self.root))

    def test_provider_detection_never_guesses_from_installation(self):
        self.assertEqual('', detect_provider(environ={}))
        self.assertEqual('claude', detect_provider(environ={'CLAUDECODE': '1'}))
        self.assertEqual('codex', detect_provider(environ={'CODEX_THREAD_ID': 'thread'}))
        self.assertEqual('', detect_provider(environ={'CLAUDECODE': '1', 'CODEX_THREAD_ID': 'thread'}))
        self.assertEqual('claude', detect_provider('claude', {'CODEX_THREAD_ID': 'thread'}))
        with self.assertRaises(ReviewError):
            detect_provider('unknown')

    def test_missing_or_unknown_provider_is_not_silently_replaced(self):
        app = ReviewApplication(self.root, provider='codex')
        for provider in ('', 'unknown', 'claude', 'cursor', 'copilot'):
            with patch('review_app.shutil.which', return_value=None), self.assertRaises(ReviewError):
                app.submit(upload(provider=provider))
        self.assertEqual([], list(self.root.iterdir()))
        self.assertFalse(app.busy)

    def test_dispatch_uses_selected_provider_and_saves_it(self):
        for provider in ('codex', 'claude', 'cursor', 'copilot'):
            with self.subTest(provider=provider), patch('review_app.shutil.which', return_value='/bin/tool'), \
                    patch('review_app.run_codex', return_value=valid_review()) as codex, \
                    patch('review_app.run_claude', return_value=valid_review()) as claude, \
                    patch('review_app.run_cursor', return_value=valid_review()) as cursor, \
                    patch('review_app.run_copilot', return_value=valid_review()) as copilot:
                app = ReviewApplication(self.root, provider='codex')
                identifier = app.submit(upload(provider=provider))
                status = self.wait_job(app, identifier)
                self.assertEqual('complete', status['status'])
                self.assertEqual(provider, status['provider'])
                self.assertEqual(int(provider == 'codex'), codex.call_count)
                self.assertEqual(int(provider == 'claude'), claude.call_count)
                self.assertEqual(int(provider == 'cursor'), cursor.call_count)
                self.assertEqual(int(provider == 'copilot'), copilot.call_count)
                saved = json.loads((app.directory(identifier) / 'request.json').read_text())
                self.assertEqual(provider, saved['provider'])

    def test_claude_failure_does_not_fall_back_to_codex(self):
        with patch('review_app.shutil.which', return_value='/bin/tool'), \
                patch('review_app.run_claude', side_effect=ReviewError('Claude sign-in required')), \
                patch('review_app.run_codex') as codex:
            app = ReviewApplication(self.root, provider='claude')
            status = self.wait_job(app, app.submit(upload(provider='claude')))
            self.assertEqual('failed', status['status'])
            self.assertIn('Claude', status['message'])
            codex.assert_not_called()

    def test_claude_structured_output_and_restricted_tools(self):
        def fake_run(command, **kwargs):
            self.assertEqual('json', command[command.index('--output-format') + 1])
            self.assertEqual('WebFetch,WebSearch', command[command.index('--tools') + 1])
            self.assertIn('--strict-mcp-config', command)
            self.assertIn('--disable-slash-commands', command)
            self.assertNotIn('--dangerously-skip-permissions', command)
            self.assertEqual(self.root, kwargs['cwd'])
            self.assertNotIn('CLAUDECODE', kwargs['env'])
            self.assertEqual('skill prompt', kwargs['input'])
            self.assertTrue(json.loads(command[command.index('--settings') + 1])['disableAllHooks'])
            return type('Result', (), {'returncode': 0, 'stdout': json.dumps({
                'is_error': False, 'structured_output': {'review': valid_review(), 'error': ''}})})()
        with patch('review_app.shutil.which', return_value='/bin/claude'), \
                patch('review_app.subprocess.run', side_effect=fake_run), \
                patch.dict('os.environ', {'CLAUDECODE': '1'}):
            self.assertEqual(valid_review(), run_claude('skill prompt', self.root))

    def test_claude_errors_are_actionable_without_exposing_output(self):
        for output in ('private secret log', json.dumps({'is_error': True, 'result': 'private secret'}),
                       json.dumps({'structured_output': {'review': None, 'error': 'private secret'}})):
            process = type('Result', (), {'returncode': 0, 'stdout': output})()
            with patch('review_app.shutil.which', return_value='/bin/claude'), \
                    patch('review_app.subprocess.run', return_value=process), self.assertRaises(ReviewError) as caught:
                run_claude('prompt', self.root)
            self.assertNotIn('private secret', str(caught.exception))

    def test_cursor_executable_aliases_and_setup_display(self):
        with patch('review_app.shutil.which', side_effect=lambda name: '/bin/agent' if name == 'agent' else None):
            self.assertEqual('/bin/agent', provider_executable('cursor'))
            config = provider_config('cursor')
            cursor = next(item for item in config['providers'] if item['id'] == 'cursor')
            self.assertTrue(cursor['installed'])
            self.assertIn('agent login', cursor['setup'])

    def test_cursor_and_copilot_output_and_permissions(self):
        for provider, runner in (('cursor', run_cursor), ('copilot', run_copilot)):
            def fake_run(command, **kwargs):
                self.assertNotIn('private resume prompt', command)
                self.assertIn('private resume prompt', kwargs['input'])
                self.assertIn('"evidence_ids"', kwargs['input'])
                self.assertEqual(self.root, kwargs['cwd'])
                self.assertEqual(300, kwargs['timeout'])
                self.assertNotIn('--force', command)
                self.assertNotIn('--allow-all', command)
                if provider == 'cursor':
                    self.assertEqual('ask', command[command.index('--mode') + 1])
                    permissions = json.loads((self.root/'.cursor/cli.json').read_text())['permissions']
                    for restriction in ('Shell(*)', 'Write(**)', 'Mcp(*:*)'):
                        self.assertIn(restriction, permissions['deny'])
                else:
                    self.assertIn('--available-tools=web_fetch', command)
                    self.assertIn('--deny-tool=shell,write,read', command)
                    self.assertIn('--no-remote-export', command)
                output = json.dumps({'review': valid_review(), 'error': ''})
                if provider == 'cursor':
                    output = json.dumps({'type': 'result', 'subtype': 'success', 'is_error': False, 'result': output})
                return type('Result', (), {'returncode': 0, 'stdout': output})()
            with self.subTest(provider=provider), patch('review_app.shutil.which', return_value='/bin/ai'), \
                    patch('review_app.subprocess.run', side_effect=fake_run):
                self.assertEqual(valid_review(), runner('private resume prompt', self.root))

    def test_text_provider_invalid_outputs_fail_without_exposing_logs(self):
        invalid_review = valid_review()
        invalid_review['bullet_reviews'][0]['rating'] = 'made-up-rating'
        for provider, runner in (('cursor', run_cursor), ('copilot', run_copilot)):
            for output in ('private secret logs', json.dumps({'review': invalid_review, 'error': ''}),
                           json.dumps({'review': valid_review()}), json.dumps({'review': [], 'error': ''})):
                if provider == 'cursor':
                    output = json.dumps({'type': 'result', 'is_error': False, 'result': output})
                process = type('Result', (), {'returncode': 0, 'stdout': output})()
                with patch('review_app.shutil.which', return_value='/bin/ai'), \
                        patch('review_app.subprocess.run', return_value=process), self.assertRaises(ReviewError) as caught:
                    runner('prompt', self.root)
                self.assertNotIn('private secret', str(caught.exception))

    def test_cursor_and_copilot_missing_auth_failure_and_timeout(self):
        import subprocess
        for provider, runner in (('cursor', run_cursor), ('copilot', run_copilot)):
            with patch('review_app.shutil.which', return_value=None), self.assertRaises(ReviewError):
                runner('prompt', self.root)
            process = type('Result', (), {'returncode': 1, 'stdout': 'private authentication details'})()
            with patch('review_app.shutil.which', return_value='/bin/ai'), \
                    patch('review_app.subprocess.run', return_value=process), self.assertRaises(ReviewError) as caught:
                runner('prompt', self.root)
            self.assertIn('sign in' if provider == 'copilot' else 'login', str(caught.exception))
            self.assertNotIn('private authentication', str(caught.exception))
            with patch('review_app.shutil.which', return_value='/bin/ai'), \
                    patch('review_app.subprocess.run', side_effect=subprocess.TimeoutExpired('ai', 300)), \
                    self.assertRaisesRegex(ReviewError, 'took too long'):
                runner('prompt', self.root)

    def test_http_submission_status_result_and_change_job(self):
        server = create_app_server(self.root, runner=lambda *args: valid_review())
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        origin = f'http://127.0.0.1:{server.server_port}'
        try:
            with self.assertRaises(urllib.error.HTTPError) as host_error:
                urllib.request.urlopen(urllib.request.Request(origin, headers={'Host': 'foreign.example'}))
            self.assertEqual(403, host_error.exception.code)
            with urllib.request.urlopen(origin) as response:
                html = response.read().decode()
            token = re.search('name="review-token" content="([^"]+)"', html)[1]
            self.assertIn('Drag your resume here', html)
            def post(path, payload, request_origin=origin):
                request = urllib.request.Request(origin + path, data=json.dumps(payload).encode(),
                    headers={'Origin': request_origin, 'X-Review-Token': token, 'Content-Type': 'application/json'})
                with urllib.request.urlopen(request) as response:
                    self.assertEqual(202, response.status)
                    return json.loads(response.read())
            with self.assertRaises(urllib.error.HTTPError) as error:
                post('/api/reviews', upload(), 'https://foreign.example')
            self.assertEqual(403, error.exception.code)
            result = post('/api/reviews', upload(provider='copilot'))
            status = self.wait_job(server.application, result['id'])
            with urllib.request.urlopen(origin + '/api/reviews/' + result['id']) as response:
                self.assertEqual('complete', json.loads(response.read())['status'])
            with urllib.request.urlopen(origin + status['review_url']) as response:
                html = response.read().decode()
            self.assertIn('Review another resume', html)
            self.assertIn('/api/reviews/' + result['id'] + '/job', html)
            changed = post('/api/reviews/' + result['id'] + '/job', {'source': SOURCE})
            self.assertNotEqual(result['id'], changed['id'])
            changed_status = self.wait_job(server.application, changed['id'])
            self.assertEqual('complete', changed_status['status'])
            self.assertEqual('copilot', changed_status['provider'])
            with self.assertRaises(urllib.error.HTTPError):
                urllib.request.urlopen(origin + '/reviews/../../README.md')
        finally:
            server.shutdown()
            server.server_close()
            thread.join()


if __name__ == '__main__':
    unittest.main()
